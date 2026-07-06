"""File upload + chunking for chat-scoped document interaction.

Uploaded files are stored in MongoDB, chunked into segments, and made
available to the agent via the `document_search` tool.
"""

from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Any

from app.core.mongo import mongo_db

COLLECTION = "chat_documents"
CHUNK_SIZE = 1500  # chars per chunk
CHUNK_OVERLAP = 200


async def ensure_indexes() -> None:
    try:
        await mongo_db[COLLECTION].create_index("chat_id")
        await mongo_db[COLLECTION].create_index([("chat_id", 1), ("filename", 1)])
    except Exception:
        pass


def _chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk.strip())
        start = end - overlap
    return chunks


def _extract_text(content: bytes, filename: str) -> str:
    """Extract text from various file formats."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("txt", "md", "csv", "json", "xml", "html", "log"):
        return content.decode("utf-8", errors="replace")

    if ext == "pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=content, filetype="pdf")
            text = ""
            for page in doc:
                text += page.get_text() + "\n"
            doc.close()
            return text
        except ImportError:
            # Fallback: try pdfplumber
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(content)) as pdf:
                    return "\n".join(page.extract_text() or "" for page in pdf.pages)
            except ImportError:
                return f"[PDF file: {filename} — install PyMuPDF or pdfplumber for text extraction]"

    if ext in ("docx",):
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            return f"[DOCX file: {filename} — install python-docx for text extraction]"

    if ext in ("xlsx", "xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
            return "\n".join(rows)
        except Exception:
            return f"[Excel file: {filename}]"

    # Unknown format — try as text
    try:
        return content.decode("utf-8", errors="replace")
    except Exception:
        return f"[Binary file: {filename} — {len(content)} bytes]"


async def upload_document(
    chat_id: str,
    filename: str,
    content: bytes,
    user_email: str,
) -> dict[str, Any]:
    """Upload a file, extract text, chunk it, and store in MongoDB."""
    text = _extract_text(content, filename)
    chunks = _chunk_text(text)
    file_hash = hashlib.sha256(content).hexdigest()[:16]

    doc = {
        "chat_id": chat_id,
        "filename": filename,
        "file_hash": file_hash,
        "file_size": len(content),
        "text_length": len(text),
        "chunk_count": len(chunks),
        "chunks": chunks,
        "uploaded_by": user_email,
        "uploaded_at": datetime.now(timezone.utc),
    }
    result = await mongo_db[COLLECTION].insert_one(doc)
    return {
        "id": str(result.inserted_id),
        "filename": filename,
        "file_size": len(content),
        "text_length": len(text),
        "chunk_count": len(chunks),
        "preview": text[:300],
    }


async def search_documents(chat_id: str, query: str, top_k: int = 5) -> list[dict]:
    """Search uploaded documents for a chat by text matching."""
    docs = await mongo_db[COLLECTION].find(
        {"chat_id": chat_id}, {"_id": 0, "filename": 1, "chunks": 1}
    ).to_list(length=20)

    results = []
    q_lower = query.lower()
    for doc in docs:
        for i, chunk in enumerate(doc.get("chunks", [])):
            if q_lower in chunk.lower():
                results.append({
                    "filename": doc["filename"],
                    "chunk_index": i,
                    "text": chunk[:500],
                    "relevance": "match",
                })
                if len(results) >= top_k:
                    return results
    # If no exact matches, return first chunks from all docs
    if not results:
        for doc in docs:
            if doc.get("chunks"):
                results.append({
                    "filename": doc["filename"],
                    "chunk_index": 0,
                    "text": doc["chunks"][0][:500],
                    "relevance": "first_chunk",
                })
    return results[:top_k]


async def list_documents(chat_id: str) -> list[dict]:
    """List all uploaded documents for a chat."""
    docs = await mongo_db[COLLECTION].find(
        {"chat_id": chat_id},
        {"_id": 0, "chunks": 0},  # exclude chunks for list view
    ).sort("uploaded_at", -1).to_list(length=50)
    for d in docs:
        if "uploaded_at" in d and hasattr(d["uploaded_at"], "isoformat"):
            d["uploaded_at"] = d["uploaded_at"].isoformat()
    return docs


async def get_context_info(chat_id: str) -> dict:
    """Return context metadata for the UI indicator."""
    from app.memory.short_term import get_short_term_memory, get_conversation_summary

    messages = await get_short_term_memory(chat_id)
    summary = await get_conversation_summary(chat_id)
    docs = await mongo_db[COLLECTION].find(
        {"chat_id": chat_id}, {"_id": 0, "filename": 1, "chunk_count": 1, "text_length": 1}
    ).to_list(length=20)

    total_doc_chunks = sum(d.get("chunk_count", 0) for d in docs)
    total_doc_chars = sum(d.get("text_length", 0) for d in docs)
    msg_chars = sum(len(m.get("content", "")) for m in messages)
    summary_chars = len(summary)

    # Rough token estimate (1 token ≈ 4 chars)
    est_tokens = (msg_chars + summary_chars + min(total_doc_chars, 6000)) // 4

    return {
        "message_count": len(messages),
        "has_summary": bool(summary),
        "summary_chars": summary_chars,
        "document_count": len(docs),
        "document_chunks": total_doc_chunks,
        "estimated_tokens": est_tokens,
        "documents": [{"filename": d["filename"], "chunks": d.get("chunk_count", 0)} for d in docs],
    }
